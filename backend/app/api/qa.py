from __future__ import annotations

import csv
import io
import re

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import any_staff
from app.db.session import get_session
from app.llm.providers import LLMError
from app.models import QAPair, UnansweredQuestion, Variable
from app.rag.documents import bump_context_version
from app.core.config import get_settings
from app.rag.embeddings import HashingEmbedder, embed_texts, embedding_signature
from app.rag.extract import _pick
from app.rag.templates import normalize_question
from app.rag.search import content_words
from app.rag.categories import CATEGORIES, KNOWN_NAMES, guess_category, standard_category
from app.services import learning
from app.services.answer_cache import forget_value
from app.services.settings_service import get_setting

router = APIRouter(prefix="/api", tags=["qa"], dependencies=[Depends(any_staff)])


class QAIn(BaseModel):
    question: str = Field(min_length=1, max_length=1000)
    answer_template: str = Field(min_length=1, max_length=4000)
    category: str | None = Field(None, max_length=64)
    data_class: str = Field("fixed", pattern="^(fixed|variable)$")
    active: bool = True
    resolve_unanswered_id: int | None = None   # "answer is wrong"/unanswered -> corrected answer


def _category(raw: str | None) -> str | None:
    return standard_category(raw)


def _row(q: QAPair) -> dict:
    return {"id": q.id, "question": q.question, "answer_template": q.answer_template, "category": q.category,
            "data_class": q.data_class, "active": q.active, "updated_at": q.updated_at,
            "keywords": list(dict.fromkeys(content_words(q.question)))[:12], "keyword_path": [(q.category or "Uncategorized"), "question", str(q.id)]}


def _unknown_row(u: UnansweredQuestion) -> dict:
    return {"id": u.id, "question": u.question, "channel": u.channel, "reason": u.reason, "status": u.status,
            "occurrences": u.occurrences, "last_seen_at": u.last_seen_at, "suggested_category": u.suggested_category,
            "suggested_keywords": u.suggested_keywords or [], "draft_answer": u.draft_answer, "wrong_answer": u.wrong_answer,
            "conversation_id": u.conversation_id, "created_at": u.created_at}


@router.get("/unanswered")
async def list_unanswered(limit: int = 100, status: str | None = None, session: AsyncSession = Depends(get_session)):
    """The approval queue: open items (new / reviewing) by default, most asked first. status=approved|ignored|all."""
    q = select(UnansweredQuestion)
    if status in learning.STATUSES:
        q = q.where(UnansweredQuestion.status == status)
    elif status != "all":
        q = q.where(UnansweredQuestion.status.in_(learning.OPEN))
    rows = (await session.execute(q.order_by(UnansweredQuestion.occurrences.desc(),
                                             UnansweredQuestion.last_seen_at.desc().nullslast(), UnansweredQuestion.id.desc())
                                  .limit(min(limit, 500)))).scalars().all()
    return [_unknown_row(u) for u in rows]


class UnknownIn(BaseModel):
    status: str = Field(pattern="^(new|reviewing|ignored)$")  # "approved" only by saving the Q&A ("Add answer")


@router.patch("/unanswered/{uid}")
async def set_unknown_status(uid: int, body: UnknownIn, session: AsyncSession = Depends(get_session)):
    u = await session.get(UnansweredQuestion, uid)
    if not u:
        raise HTTPException(404, "Not found")
    learning.set_status(u, body.status)
    await session.commit()
    return _unknown_row(u)


@router.post("/unanswered/{uid}/draft")
async def draft_unknown(uid: int, session: AsyncSession = Depends(get_session)):
    """Suggest an answer from the existing knowledge. Saved on the row as a draft only; nothing goes live."""
    from app.services.reply_engine import get_llm_service
    u = await session.get(UnansweredQuestion, uid)
    if not u:
        raise HTTPException(404, "Not found")
    u.draft_answer = await learning.draft_for(session, get_llm_service(session), u)
    if u.status == "new":
        learning.set_status(u, "reviewing")
    await session.commit()
    return _unknown_row(u)


async def _embed(session, texts):
    try:
        return await embed_texts(session, texts)
    except LLMError as exc:
        raise HTTPException(502, f"Embedding failed: {exc}") from exc


@router.get("/qa")
async def list_qa(search: str | None = None, category: str | None = None, data_kind: str | None = None,
                  limit: int = 100, offset: int = 0,
                  session: AsyncSession = Depends(get_session)):
    q = select(QAPair)
    if search:
        q = q.where(or_(QAPair.question.ilike(f"%{search}%"), QAPair.answer_template.ilike(f"%{search}%")))
    if category:
        q = q.where(QAPair.category == category)
    if data_kind == "fixed":
        q = q.where(QAPair.data_class == "fixed")
    elif data_kind == "variable":
        q = q.where(QAPair.data_class == "variable")
    total = (await session.execute(select(func.count()).select_from(q.subquery()))).scalar_one()
    rows = (await session.execute(q.order_by(QAPair.id.desc()).limit(min(limit, 500)).offset(offset))).scalars().all()
    used = (await session.execute(select(QAPair.category).where(QAPair.category.is_not(None)).distinct())).scalars().all()
    cats = list(CATEGORIES) + sorted(c for c in used if c not in CATEGORIES)
    return {"total": total, "items": [_row(r) for r in rows], "categories": cats}


@router.post("/qa")
async def create_qa(body: QAIn, session: AsyncSession = Depends(get_session)):
    vec = (await _embed(session, [body.question]))[0]
    version = await embedding_signature(session)
    qa = QAPair(question=body.question.strip(), question_norm=normalize_question(body.question), answer_template=body.answer_template,
                category=_category(body.category) or guess_category(body.question), data_class=body.data_class, active=body.active,
                embedding=vec, embedding_version=version)
    session.add(qa)
    if body.resolve_unanswered_id:  # the admin approved the (edited) draft: this is the only way a draft goes live
        u = await session.get(UnansweredQuestion, body.resolve_unanswered_id)
        if u:
            learning.set_status(u, "approved")
    await bump_context_version(session)
    await session.commit()
    await session.refresh(qa)
    return _row(qa)


@router.put("/qa/{qid}")
async def update_qa(qid: int, body: QAIn, session: AsyncSession = Depends(get_session)):
    qa = await session.get(QAPair, qid)
    if not qa:
        raise HTTPException(404, "Not found")
    if body.question.strip() != qa.question:
        qa.embedding = (await _embed(session, [body.question]))[0]
        qa.embedding_version = await embedding_signature(session)
    qa.question, qa.question_norm = body.question.strip(), normalize_question(body.question)
    qa.answer_template, qa.category, qa.data_class, qa.active = body.answer_template, _category(body.category), body.data_class, body.active
    await bump_context_version(session)
    await session.commit()
    await session.refresh(qa)
    return _row(qa)


@router.delete("/qa/{qid}")
async def delete_qa(qid: int, session: AsyncSession = Depends(get_session)):
    qa = await session.get(QAPair, qid)
    if not qa:
        raise HTTPException(404, "Not found")
    await session.delete(qa)
    await bump_context_version(session)
    await session.commit()
    return {"ok": True}


def _safe_cell(v: str) -> str:
    """Neutralize spreadsheet formula injection in exported CSV."""
    return "'" + v if v[:1] in ("=", "+", "-", "@", "\t", "\r") else v


@router.get("/qa-export")
async def export_qa(session: AsyncSession = Depends(get_session)):
    """Every Q&A as stored. The first four columns are the import format (edit and import the file back); the rest
    show how the assistant keeps each row: its matching form, keyword-index terms, variables, embedding state."""
    from app.models import KeywordEntry  # local: only the export needs it
    current = await embedding_signature(session)
    variables = {v.key: v.value for v in (await session.execute(select(Variable))).scalars()}
    terms: dict[str, list[str]] = {}
    for ref, term in (await session.execute(select(KeywordEntry.source_ref, KeywordEntry.term).where(
            KeywordEntry.source_type == "qa").order_by(KeywordEntry.weight.desc()))).all():
        if term not in terms.setdefault(ref, []):
            terms[ref].append(term)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["question", "answer_template", "category", "active", "id", "matched_as", "keywords", "variables",
                "missing_variables", "answer_preview", "embedding", "updated_at"])
    for q in (await session.execute(select(QAPair).order_by(QAPair.category.nullsfirst(), QAPair.id))).scalars():
        used = list(dict.fromkeys(re.findall(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}", q.answer_template)))
        missing = [v for v in used if not (variables.get(v) or "").strip()]
        preview = re.sub(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}", lambda m: variables.get(m.group(1)) or m.group(0), q.answer_template)
        w.writerow([_safe_cell(q.question), _safe_cell(q.answer_template), _safe_cell(q.category or ""), str(q.active).lower(),
                    q.id, _safe_cell(q.question_norm), _safe_cell(", ".join(terms.get(str(q.id), [])[:15])),
                    ", ".join(used), ", ".join(missing),
                    _safe_cell(preview if not missing else "(not used until these variables have values)"),
                    "ready" if q.embedding_version == current else "pending (text matching works meanwhile)",
                    q.updated_at.isoformat() if q.updated_at else ""])
    return Response("\ufeff" + buf.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"content-disposition": 'attachment; filename="qa_pairs.csv"'})


def _read_records(data: bytes, ext: str) -> list[dict]:
    """Rows of a Q&A file: CSV/TSV, JSON (list of objects, {"items": [...]}, or JSON lines) or Excel."""
    import json
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        head = [str(h or "").strip() for h in (rows[0] if rows else [])]
        return [dict(zip(head, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    text = data.decode("utf-8-sig", errors="replace")
    if ext in ("json", "jsonl") or text.lstrip()[:1] in ("[", "{"):
        try:
            d = json.loads(text)
        except ValueError:
            try:
                d = [json.loads(line) for line in text.splitlines() if line.strip()]
            except ValueError:
                raise HTTPException(422, "Could not read the JSON file") from None
        if isinstance(d, dict):
            d = next((v for v in d.values() if isinstance(v, list)), [d])
        return [x for x in d if isinstance(x, dict)]
    delim = "\t" if ext == "tsv" else ","
    return list(csv.DictReader(io.StringIO(text), delimiter=delim))


async def _embed_or_defer(session, texts: list[str]) -> tuple[list, list[str], bool]:
    """Embed in batches. When the provider refuses part way (free-tier quota: 429), the rest get a local stand-in
    vector marked with the hash signature; the re-embedding job replaces them later. Saved Q&A text matching works
    either way, so an import never fails because of the embedding quota."""
    vecs: list = []
    versions: list[str] = []
    real = await embedding_signature(session)
    for i in range(0, len(texts), 64):  # embedding APIs cap the batch size
        batch = texts[i:i + 64]
        try:
            vecs += await embed_texts(session, batch)
            versions += [real] * len(batch)
        except LLMError:
            local = HashingEmbedder(get_settings().embedding_dim)
            rest = texts[i:]
            vecs += await embed_texts(session, rest, embedder=(local, None))
            versions += [await embedding_signature(session, embedder=local)] * len(rest)
            return vecs, versions, True
    return vecs, versions, False


@router.post("/qa-import")
async def import_qa(request: Request, file: UploadFile = File(...), session: AsyncSession = Depends(get_session)):
    data = await file.read(10 * 1024 * 1024)
    records = _read_records(data, (file.filename or "").rsplit(".", 1)[-1].lower())
    keys = list({k for r in records for k in r})
    cols = {c.strip().lower(): c for c in keys}
    qk = _pick(keys, ("question", "questions", "q", "query", "prosno", "proshno", "প্রশ্ন"))
    ak = cols.get("answer_template") or _pick(keys, ("answer", "answers", "a", "reply", "response", "uttor", "উত্তর"))
    if not qk or not ak:
        raise HTTPException(422, "The file needs a question column/key and an answer column/key "
                                 "(e.g. 'question' and 'answer'). CSV, JSON and Excel are supported.")
    rows = [{k: ("" if v is None else str(v)) for k, v in r.items()} for r in records]
    rows = [r for r in rows if r.get(qk, "").strip() and r.get(ak, "").strip()][:2000]
    # merge: a question already saved (same wording) gets the file's answer/category/active; its embedding (of the
    # question) stays. Only new questions are embedded.
    existing = {q.question_norm: q for q in (await session.execute(select(QAPair))).scalars()}
    catk = cols.get("category") or cols.get("topic")
    new_rows: dict[str, dict] = {}
    created = updated = unchanged = 0
    for r in rows:
        norm = normalize_question(r[qk])
        active = (r.get(cols.get("active", "")) or "true").strip().lower() not in ("false", "0", "no")
        raw = (r.get(catk) or "").strip()[:64] if catk else ""
        # no category in the file: an existing Q&A keeps its own, a new one gets a guess from its question
        cat = (_category(raw) if raw.lower() in KNOWN_NAMES else guess_category(r[qk])) if raw else None
        answer = r[ak].strip()
        if norm in existing:
            q = existing[norm]
            cat = cat or q.category or guess_category(r[qk])
            if (q.answer_template, q.category, q.active) == (answer, cat, active):
                unchanged += 1
            else:
                q.answer_template, q.category, q.active = answer, cat, active
                updated += 1
        elif norm in new_rows:  # the same question twice in the file: the later row wins
            new_rows[norm].update(answer=answer, category=cat or new_rows[norm]["category"], active=active)
        else:
            new_rows[norm] = {"question": r[qk].strip(), "answer": answer, "category": cat or guess_category(r[qk]),
                              "active": active}
    vecs, versions, deferred = await _embed_or_defer(session, [n["question"] for n in new_rows.values()])
    for (norm, n), vec, version in zip(new_rows.items(), vecs, versions):
        session.add(QAPair(question=n["question"], question_norm=norm, answer_template=n["answer"], category=n["category"],
                           active=n["active"], embedding=vec, embedding_version=version))
        created += 1
    await bump_context_version(session)
    await session.commit()
    queued = False
    if deferred and (queue := getattr(request.app.state, "queue", None)) is not None:
        queued = bool(await queue.enqueue_job("reembed_all_job", _job_id=f"reembed-qa-import-{created}-{len(rows)}"))
    variables = {v for (v,) in (await session.execute(select(Variable.key))).all()}
    used = sorted({m for n in rows for m in re.findall(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}", n[ak])})
    return {"created": created, "updated": updated, "unchanged": unchanged, "skipped": len(records) - len(rows),
            "embedding_deferred": deferred, "reembed_queued": queued,
            "missing_variables": [v for v in used if v not in variables]}


class VarIn(BaseModel):
    value: str = Field(max_length=2000)
    description: str | None = Field(None, max_length=300)  # how guests ask for it; None = keep the current one


@router.get("/variables")
async def list_vars(session: AsyncSession = Depends(get_session)):
    return [{"key": v.key, "value": v.value, "description": v.description}
            for v in (await session.execute(select(Variable).order_by(Variable.key))).scalars()]


@router.get("/variables/usage")
async def vars_usage(session: AsyncSession = Depends(get_session)):
    """Every {variable} the active Q&A answers use: how many answers and a few of their questions. A key with no value
    here (or an empty one) keeps all those answers from being sent, so the Variables tab lists it to be filled in."""
    used: dict[str, dict] = {}
    rows = await session.execute(select(QAPair.question, QAPair.answer_template).where(QAPair.active.is_(True)).order_by(QAPair.id))
    for question, answer in rows:
        for key in dict.fromkeys(re.findall(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}", answer)):
            u = used.setdefault(key, {"count": 0, "examples": []})
            u["count"] += 1
            if len(u["examples"]) < 3:
                u["examples"].append(question)
    return used


@router.put("/variables/{key}")
async def put_var(key: str, body: VarIn, session: AsyncSession = Depends(get_session)):
    import re
    if not re.fullmatch(r"[A-Za-z0-9_]{1,64}", key):
        raise HTTPException(422, "Key may contain letters, digits and underscores only")
    v = await session.get(Variable, key)
    if v is None:
        v = Variable(key=key, value=body.value)
        session.add(v)
    elif v.value != body.value:
        await forget_value(session, v.value)  # cached answers that quote the old value are wrong now
    v.value = body.value
    if body.description is not None:
        v.description = body.description.strip() or None
    await bump_context_version(session)
    await session.commit()
    return {"key": key, "value": v.value, "description": v.description}


@router.delete("/variables/{key}")
async def del_var(key: str, session: AsyncSession = Depends(get_session)):
    v = await session.get(Variable, key)
    if v:
        await forget_value(session, v.value)
        await session.delete(v)
        await bump_context_version(session)
        await session.commit()
    return {"ok": True}
