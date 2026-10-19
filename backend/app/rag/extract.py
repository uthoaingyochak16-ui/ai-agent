"""Text extraction for (almost) any uploaded file. Returns [(page_no | None, text), ...].

Native parsers:  pdf, docx, xlsx/xlsm, pptx, csv/tsv, html/xml, rtf, odt/ods/odp, epub, zip, plain text of any kind
External tools:  images and scanned PDF pages -> tesseract OCR (English + Bangla), via pdftoppm for PDFs
                 legacy / other office formats (doc, xls, ppt, pages, wpd...) -> LibreOffice -> PDF
Anything else is accepted when it decodes as text; truly binary media (audio, video) fails with a clear error.
"""
from __future__ import annotations

import csv
import io
import logging
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

logger = logging.getLogger(__name__)

TEXT_TYPES = {"txt", "text", "md", "markdown", "rst", "log", "yaml", "yml", "ini", "cfg", "conf",
              "toml", "srt", "vtt", "sql", "tex", "properties", "env"}
HTML_TYPES = {"html", "htm", "xhtml", "xml", "svg"}
IMAGE_TYPES = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff", "webp", "pbm", "pgm", "ppm"}
OFFICE_TYPES = {"doc", "dot", "docm", "dotx", "xls", "xlt", "xlsb", "ppt", "pps", "ppsx", "pptm", "pot", "potx",
                "odg", "fodt", "fods", "fodp", "wpd", "wps", "pages", "key", "numbers", "abw", "hwp", "pub", "vsd", "vsdx"}
MEDIA_TYPES = {"mp3", "wav", "ogg", "m4a", "aac", "flac", "opus", "mp4", "mov", "avi", "mkv", "webm", "3gp"}

OCR_LANGS = "eng+ben"
OCR_MAX_PAGES = 60
OCR_MIN_CHARS = 25           # a PDF page with less text than this is treated as scanned
TOOL_TIMEOUT_S = 240
ZIP_MAX_FILES, ZIP_MAX_BYTES = 200, 200 * 1024 * 1024


class ExtractionError(ValueError):
    pass


# ---- helpers -----------------------------------------------------------------------------------------
def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _looks_text(data: bytes) -> bool:
    head = data[:8192]
    if not head or b"\x00" in head and not head.startswith((b"\xff\xfe", b"\xfe\xff")):
        return False
    try:
        head.decode("utf-8")
        return True
    except UnicodeDecodeError as e:  # a multi-byte character cut at the 8 KiB boundary is fine
        return e.start >= len(head) - 4


class _TextOnly(HTMLParser):
    _SKIP = {"script", "style", "noscript", "head"}
    _BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "td", "th"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip += 1
        elif tag in self._BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skip:
            self._skip -= 1

    def handle_data(self, d):
        if not self._skip:
            self.out.append(d)


def _strip_markup(text: str) -> str:
    p = _TextOnly()
    p.feed(text)
    return re.sub(r"\n\s*\n+", "\n", "".join(p.out)).strip()


def _need(tool: str, what: str) -> str:
    path = shutil.which(tool)
    if not path:
        raise ExtractionError(f"{what} needs '{tool}', which is not installed on the server")
    return path


def _run(cmd: list[str]) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, capture_output=True, timeout=TOOL_TIMEOUT_S, check=False)


def sniff(data: bytes) -> str:
    """Guess a type from the content when the extension is missing or unknown."""
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"\x89PNG"):
        return "png"
    if data.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if data.startswith((b"GIF87a", b"GIF89a")):
        return "gif"
    if data.startswith((b"II*\x00", b"MM\x00*")):
        return "tiff"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    if data.startswith(b"{\\rtf"):
        return "rtf"
    if data.startswith(b"\xd0\xcf\x11\xe0"):  # old MS Office container (doc/xls/ppt)
        return "doc"
    if data.startswith(b"PK\x03\x04"):
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile:
            return "zip"
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "ppt/presentation.xml" in names:
            return "pptx"
        if "content.xml" in names:
            return "odt"
        if "META-INF/container.xml" in names:
            return "epub"
        return "zip"
    if _looks_text(data):
        head = data[:512].lstrip().lower()
        return "html" if head.startswith((b"<!doctype html", b"<html")) else "txt"
    return "bin"


# ---- per-format extractors -----------------------------------------------------------------------------
# Column / key names that mark a Q&A-shaped record (any language the files are likely to use)
_Q_KEYS = ("question", "questions", "q", "query", "prosno", "proshno", "prashna", "প্রশ্ন", "title", "heading", "topic_question")
_A_KEYS = ("answer", "answers", "short_answer", "long_answer", "detailed_answer", "full_answer", "a", "answer_template",
           "reply", "response", "uttor", "uttar", "উত্তর", "content", "body", "text",
           "description", "details")


# Fields that describe what a record is about: not part of the answer, but they go into the keyword index
_INDEX_META_KEYS = {"topic", "category", "categories", "tags", "keywords", "keyword"}

# Bookkeeping fields of a record: never part of the text a guest is shown
_META_KEYS = {"id", "source", "sources", "url", "link", "links", "ref", "reference", "references", "image", "images",
              "topic", "category", "categories", "tags", "created_at", "updated_at", "question_type", "type",
              "keywords", "details_ref", "ref_id", "lang", "language"}


def _pick(keys: list[str], wanted: tuple[str, ...]) -> str | None:
    low = {k.strip().lower(): k for k in keys}
    return next((low[w] for w in wanted if w in low), None)


def _records_to_sections(records: list[dict], page_key: str | None = None) -> list[tuple]:
    """One section per record. Q&A-shaped records keep their question as the section title, so every chunk cut from
    a long answer still knows which question it answers."""
    if not records:
        return []
    keys = list({k for r in records for k in r})
    qk, ak = _pick(keys, _Q_KEYS), _pick(keys, _A_KEYS)
    if (qk or "").lower() == "q" and not ak or (ak or "").lower() == "a" and not qk:
        qk = ak = None  # a lone column called "q" or "a" is just data, not a Q&A file
    pk = page_key or _pick(keys, ("page", "page_no", "pageno", "পৃষ্ঠা"))
    out = []
    for r in records:
        q = str(r.get(qk) or "").strip() if qk else ""
        a = str(r.get(ak) or "").strip() if ak else ""
        extra = [f"{k}: {v}" for k, v in r.items() if k not in (qk, ak, pk) and str(k).strip().lower() not in _META_KEYS
                 and str(v).strip() and not isinstance(v, (dict, list))]
        body = "\n".join(x for x in (a, ", ".join(extra)) if x)
        if not body and not q:
            continue
        page = r.get(pk) if pk else None
        page = int(page) if isinstance(page, (int, float)) or (isinstance(page, str) and page.isdigit()) else None
        meta = " ".join(str(v) for k, v in r.items() if str(k).strip().lower() in _INDEX_META_KEYS and str(v).strip()
                        and not isinstance(v, dict)).strip()
        out.append((page, body or q, q or None, meta or None))
    return out


def _json(data: bytes):
    import json
    text = _decode(data).strip()
    try:
        d = json.loads(text)
    except ValueError:  # JSON lines
        try:
            d = [json.loads(line) for line in text.splitlines() if line.strip()]
        except ValueError:
            return [(None, text)]
    if isinstance(d, dict):  # {"items": [...]} or {"faq": [...]}
        lists = [v for v in d.values() if isinstance(v, list) and v and all(isinstance(x, dict) for x in v)]
        if lists:
            d = [x for lst in lists for x in lst]
        else:
            return [(None, "\n".join(f"{k}: {v}" for k, v in d.items()))]
    if isinstance(d, list) and d and all(isinstance(x, dict) for x in d):
        return _records_to_sections(d)
    return [(None, json.dumps(d, ensure_ascii=False, indent=1))]


def _csv(data: bytes, delimiter: str | None = None):
    text = _decode(data)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not rows:
        return []
    head, body = [h.strip() for h in rows[0]], rows[1:]
    if not body:
        return [(None, " | ".join(head))]
    return _records_to_sections([{h: v for h, v in zip(head, r) if h} for r in body if any(c.strip() for c in r)])


def _docx(data: bytes):
    from docx import Document as Docx
    d = Docx(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [(None, "\n".join(parts))]


def _xlsx(data: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages = []
    for i, ws in enumerate(wb.worksheets, 1):
        rows = [[("" if v is None else str(v)).strip() for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        head, body = rows[0], rows[1:]
        if body and all(head):  # header row -> one section per row ("column: value"); Q&A rows keep the question
            pages += [(i, t, q, m) for _, t, q, m in _records_to_sections([dict(zip(head, r)) for r in body])]
        else:
            pages.append((i, f"[{ws.title}]\n" + "\n".join(" | ".join(v for v in r if v) for r in rows)))
    return pages


def _pptx(data: bytes):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            parts.append(slide.notes_slide.notes_text_frame.text)
        pages.append((i, "\n".join(p for p in parts if p.strip())))
    return pages


def _odf(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("content.xml").decode("utf-8", errors="replace")
    xml = re.sub(r"<text:(p|h|line-break|tab)\b", r"\n<text:\1", xml)
    return [(None, _strip_markup(xml))]


def _epub(data: bytes):
    pages = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = sorted(n for n in z.namelist() if n.lower().endswith((".xhtml", ".html", ".htm")))
        for i, n in enumerate(names, 1):
            text = _strip_markup(z.read(n).decode("utf-8", errors="replace"))
            if text:
                pages.append((i, text))
    return pages


def _rtf(data: bytes):
    from striprtf.striprtf import rtf_to_text
    return [(None, rtf_to_text(_decode(data), errors="ignore"))]


def _ocr_image_file(path: Path) -> str:
    tess = _need("tesseract", "Reading text from images (OCR)")
    r = _run([tess, str(path), "stdout", "-l", OCR_LANGS])
    if r.returncode != 0:
        raise ExtractionError(f"OCR failed: {r.stderr.decode(errors='replace')[:200]}")
    return r.stdout.decode("utf-8", errors="replace")


def _image(data: bytes):
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "img"
        p.write_bytes(data)
        return [(None, _ocr_image_file(p))]


def _ocr_pdf_pages(data: bytes, page_numbers: list[int]) -> dict[int, str]:
    pdftoppm = _need("pdftoppm", "Reading scanned PDFs (OCR)")
    out: dict[int, str] = {}
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "in.pdf"
        src.write_bytes(data)
        for n in page_numbers[:OCR_MAX_PAGES]:
            prefix = Path(d) / f"p{n}"
            r = _run([pdftoppm, "-r", "200", "-gray", "-png", "-f", str(n), "-l", str(n), "-singlefile", str(src), str(prefix)])
            img = prefix.with_suffix(".png")
            if r.returncode == 0 and img.exists():
                out[n] = _ocr_image_file(img)
    return out


def _pdf(data: bytes):
    from pypdf import PdfReader
    pages = [(i + 1, pg.extract_text() or "") for i, pg in enumerate(PdfReader(io.BytesIO(data)).pages)]
    scanned = [n for n, t in pages if len(t.strip()) < OCR_MIN_CHARS]
    if scanned and shutil.which("tesseract") and shutil.which("pdftoppm"):
        ocr = _ocr_pdf_pages(data, scanned)
        pages = [(n, ocr.get(n) or t) for n, t in pages]
    return pages


def _via_libreoffice(data: bytes, kind: str):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ExtractionError(f".{kind} files need LibreOffice, which is not installed on the server")
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / f"in.{kind}"
        src.write_bytes(data)
        r = _run([soffice, f"-env:UserInstallation=file://{d}/profile", "--headless", "--convert-to", "pdf",
                  "--outdir", d, str(src)])
        pdf = Path(d) / "in.pdf"
        if r.returncode != 0 or not pdf.exists():
            raise ExtractionError(f"could not convert .{kind}: {r.stderr.decode(errors='replace')[:200] or 'unknown error'}")
        return _pdf(pdf.read_bytes())


def _zip(data: bytes, depth: int):
    pages = []
    total = 0
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        infos = [i for i in z.infolist() if not i.is_dir() and not Path(i.filename).name.startswith((".", "__MACOSX"))]
        for info in infos[:ZIP_MAX_FILES]:
            total += info.file_size
            if total > ZIP_MAX_BYTES:
                break
            name = info.filename
            try:
                inner = _extract(z.read(info), name.rsplit(".", 1)[-1] if "." in name else "", _depth=depth + 1)
            except Exception as exc:  # noqa: BLE001 - one unreadable member must not sink the archive
                logger.info("zip member %s skipped: %s", name, exc)
                continue
            pages += [(t[0], t[1], t[2] if len(t) > 2 else None, t[3] if len(t) > 3 else None) for t in inner if t[1].strip()]
    return pages


# ---- entry point -----------------------------------------------------------------------------------------
@dataclass(frozen=True)
class Section:
    page: int | None
    text: str
    title: str | None = None  # e.g. the question of a Q&A record
    meta: str | None = None   # topic / keywords the file gives for the record (keyword index only)


def extract_sections(data: bytes, kind: str) -> list[Section]:
    return [Section(t[0], t[1], t[2] if len(t) > 2 else None, t[3] if len(t) > 3 else None) for t in _extract(data, kind)]


def extract_pages(data: bytes, kind: str) -> list[tuple[int | None, str]]:
    """(page, text) pairs; titled sections are flattened as "title + text"."""
    return [(s.page, f"{s.title}\n{s.text}" if s.title else s.text) for s in extract_sections(data, kind)]


def _extract(data: bytes, kind: str, *, _depth: int = 0) -> list[tuple[int | None, str]]:
    kind = (kind or "").lower().lstrip(".")
    known = TEXT_TYPES | HTML_TYPES | IMAGE_TYPES | OFFICE_TYPES | {
        "pdf", "docx", "xlsx", "xlsm", "pptx", "csv", "tsv", "json", "jsonl", "rtf", "odt", "ods", "odp", "epub", "zip"}
    if kind not in known:
        if kind in MEDIA_TYPES:
            raise ExtractionError(f".{kind} is audio/video; upload a transcript (TXT/DOCX/PDF) instead")
        kind = sniff(data)
    if kind == "pdf":
        return _pdf(data)
    if kind == "docx":
        return _docx(data)
    if kind in ("xlsx", "xlsm"):
        return _xlsx(data)
    if kind == "pptx":
        return _pptx(data)
    if kind in ("json", "jsonl"):
        return _json(data)
    if kind == "csv":
        return _csv(data)
    if kind == "tsv":
        return _csv(data, "\t")
    if kind in HTML_TYPES:
        return [(None, _strip_markup(_decode(data)))]
    if kind == "rtf":
        return _rtf(data)
    if kind in ("odt", "ods", "odp"):
        return _odf(data)
    if kind == "epub":
        return _epub(data)
    if kind in IMAGE_TYPES:
        return _image(data)
    if kind == "zip":
        if _depth >= 2:
            raise ExtractionError("nested archives are too deep")
        return _zip(data, _depth)
    if kind in OFFICE_TYPES:
        return _via_libreoffice(data, kind)
    if kind in TEXT_TYPES or _looks_text(data):
        return [(None, _decode(data))]
    raise ExtractionError("no readable text in this file (unknown binary format)")
