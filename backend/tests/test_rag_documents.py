import io

import pytest
from sqlalchemy import select

from app.models import Chunk, Document
from app.rag.documents import LocalStorage, chunk_pages, extract_pages, process_document
from app.rag.embeddings import HashingEmbedder, embed_texts
from app.rag.templates import fill_template, normalize_question
from app.services.settings_service import get_setting
from tests.helpers import make_pdf

EMB = (HashingEmbedder(1024), None)


def words(n, prefix="w"):
    return " ".join(f"{prefix}{i}" for i in range(n))


def test_chunk_sizes_overlap_and_no_oversize():
    chunks = chunk_pages([(None, words(1300))])
    sizes = [len(c.split()) for _, c in chunks]
    assert all(s <= 500 for s in sizes) and sizes[0] == 400 and len(chunks) == 4
    first_tail = chunks[0][1].split()[-50:]
    assert chunks[1][1].split()[:50] == first_tail          # 50-word overlap
    assert chunk_pages([(None, "")]) == [] and len(chunk_pages([(None, words(420))])) == 1  # tiny tail absorbed


def test_chunks_keep_the_starting_page():
    pages = [(1, words(300, "a")), (2, words(300, "b"))]
    ch = chunk_pages(pages)
    assert ch[0][0] == 1 and ch[1][0] in (1, 2) and ch[-1][0] == 2


def test_bangla_text_is_chunked_by_words():
    assert len(chunk_pages([(None, " ".join(["আমাদের"] * 900))])) >= 2


def test_extractors_txt_csv_docx_pdf():
    assert "hello" in extract_pages(b"hello world", "txt")[0][1]
    csv_text = "\n".join(t for _, t in extract_pages("name,price\nTea,50\nCoffee,80\n".encode(), "csv"))  # a section per row
    assert "name: Tea, price: 50" in csv_text and "Coffee" in csv_text
    from docx import Document as D
    d = D()
    d.add_paragraph("Refunds are accepted within 7 days.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Plan", "Gold"
    buf = io.BytesIO()
    d.save(buf)
    docx_text = extract_pages(buf.getvalue(), "docx")[0][1]
    assert "Refunds" in docx_text and "Plan | Gold" in docx_text
    pdf = extract_pages(make_pdf(["First page text", "Second page text"]), "pdf")
    assert [p for p, _ in pdf] == [1, 2] and "Second page" in pdf[1][1]


def test_template_variables_and_question_normalization():
    assert fill_template("Call {phone} or {email}", {"phone": "123"}) == ("Call 123 or {email}", ["email"])
    assert normalize_question("  What're  your HOURS?? ") == "what re your hours"


def test_storage_blocks_path_traversal(tmp_path):
    st = LocalStorage(str(tmp_path))
    p = st.save("a.txt", b"x")
    assert st.read(p) == b"x"
    with pytest.raises(ValueError):
        st.read("../../etc/passwd")


async def test_hash_embedder_is_deterministic_normalized_and_similarity_ordered():
    e = HashingEmbedder(1024)
    a, b, c = await e.embed(["What are your opening hours?", "what are your opening hours", "refund policy for damaged items"])
    dot = lambda x, y: sum(i * j for i, j in zip(x, y))  # noqa: E731
    assert abs(dot(a, a) - 1) < 1e-6 and dot(a, b) > 0.95 and dot(a, c) < 0.4


async def test_external_embedders_receive_redacted_text(session):
    class Spy(HashingEmbedder):
        external = True
        seen: list[str] = []

        async def embed(self, texts):
            Spy.seen += texts
            return await super().embed(texts)
    from app.privacy.guard import ProviderPolicy
    await embed_texts(session, ["my number is 01712345678 and mail a@b.com"], embedder=(Spy(1024), ProviderPolicy(name="x")))
    assert "01712345678" not in Spy.seen[0] and "a@b.com" not in Spy.seen[0]


async def test_process_document_ready_supersedes_old_version_and_bumps_cache_version(session, tmp_path):
    st = LocalStorage(str(tmp_path))
    v0 = await get_setting(session, "context_version")
    ids = []
    for text in (b"Old price list. Tea costs 40.", b"New price list. Tea costs 50."):
        d = Document(title="Prices", file_path=st.save("p.txt", text), type="txt")
        session.add(d)
        await session.commit()
        assert await process_document(session, d.id, storage=st, embedder=EMB) == "ready"
        ids.append(d.id)
    docs = {d.id: d for d in (await session.execute(select(Document))).scalars()}
    assert docs[ids[0]].status == "superseded" and docs[ids[1]].status == "ready" and docs[ids[1]].chunk_count == 1
    chunks = (await session.execute(select(Chunk))).scalars().all()
    assert [c.document_id for c in chunks] == [ids[1]] and len(chunks[0].embedding) == 1024
    assert await get_setting(session, "context_version") > v0


async def test_empty_or_unreadable_document_is_marked_failed(session, tmp_path):
    st = LocalStorage(str(tmp_path))
    d = Document(title="Scan", file_path=st.save("s.pdf", make_pdf([""])), type="pdf")
    session.add(d)
    await session.commit()
    assert await process_document(session, d.id, storage=st, embedder=EMB) == "failed"
    await session.refresh(d)
    assert d.status == "failed" and "no readable text" in d.error
