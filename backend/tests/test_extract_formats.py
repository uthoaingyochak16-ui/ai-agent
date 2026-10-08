import io
import shutil
import zipfile

import pytest

from app.rag.extract import ExtractionError, extract_pages, sniff
from tests.helpers import make_pdf


def text_of(pages):
    return "\n".join(t for _, t in pages)


def _docx(text):
    from docx import Document
    d = Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_xlsx_sheets_become_pages_with_column_labels():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Fees"
    ws.append(["Course", "Fee"])
    ws.append(["Meditation", 1500])
    buf = io.BytesIO()
    wb.save(buf)
    pages = extract_pages(buf.getvalue(), "xlsx")
    assert pages[0][0] == 1 and "Course: Meditation, Fee: 1500" in pages[0][1]


def test_pptx_slides_and_notes():
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Admission"
    s.notes_slide.notes_text_frame.text = "Open every Friday"
    buf = io.BytesIO()
    prs.save(buf)
    t = text_of(extract_pages(buf.getvalue(), "pptx"))
    assert "Admission" in t and "Open every Friday" in t


def test_html_markdown_json_rtf_tsv():
    html = b"<html><head><style>x{}</style></head><body><h1>Hours</h1><p>9am &amp; 5pm</p><script>bad()</script></body></html>"
    t = text_of(extract_pages(html, "html"))
    assert "Hours" in t and "9am & 5pm" in t and "bad()" not in t
    assert "# Title" in text_of(extract_pages(b"# Title\nbody", "md"))
    assert "fee: 100" in text_of(extract_pages(b'{"fee": 100}', "json"))
    assert "Hello RTF" in text_of(extract_pages(rb"{\rtf1\ansi Hello RTF\par}", "rtf"))
    assert "a: 1, b: 2" in text_of(extract_pages(b"a\tb\n1\t2\n", "tsv"))


def test_bangla_text_and_unknown_text_extension():
    t = "কোর্স ফি ১৫০০ টাকা".encode()
    assert "কোর্স" in text_of(extract_pages(t, "txt"))
    assert "কোর্স" in text_of(extract_pages(t, "weird"))  # any extension that is really text


def test_odt_and_epub():
    def z(files):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as f:
            for n, c in files.items():
                f.writestr(n, c)
        return buf.getvalue()
    odt = z({"mimetype": "application/vnd.oasis.opendocument.text",
             "content.xml": '<office:document-content><text:p>Branch in Dhaka</text:p><text:p>Second line</text:p></office:document-content>'})
    assert "Branch in Dhaka" in text_of(extract_pages(odt, "odt"))
    epub = z({"META-INF/container.xml": "<container/>", "OEBPS/ch1.xhtml": "<html><body><p>Chapter one</p></body></html>"})
    assert "Chapter one" in text_of(extract_pages(epub, "epub"))
    assert sniff(odt) == "odt" and sniff(epub) == "epub"


def test_zip_reads_every_member():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as f:
        f.writestr("faq.txt", "Refund within 7 days")
        f.writestr("guide.docx", _docx("Bring your ID card"))
        f.writestr("song.mp3", b"ID3\x00\x00binary")
    t = text_of(extract_pages(buf.getvalue(), "zip"))
    assert "Refund within 7 days" in t and "Bring your ID card" in t


def test_sniff_by_content_when_extension_is_missing():
    assert sniff(make_pdf(["x"])) == "pdf"
    assert sniff(_docx("x")) == "docx"
    assert sniff(b"\x89PNG\r\n\x1a\n....") == "png"
    assert "Hello" in text_of(extract_pages(_docx("Hello"), ""))


def test_binary_and_media_fail_with_a_clear_message():
    with pytest.raises(ExtractionError, match="audio/video"):
        extract_pages(b"ID3\x00", "mp3")
    with pytest.raises(ExtractionError):
        extract_pages(b"\x00\x01\x02\x03" * 100, "bin")


@pytest.mark.skipif(not shutil.which("tesseract"), reason="tesseract not installed")
def test_image_ocr():
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("L", (900, 160), 255)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 56)
    except OSError:
        font = ImageFont.load_default(size=56)
    ImageDraw.Draw(img).text((20, 40), "Course fee 1500 Taka", fill=0, font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    assert "1500" in text_of(extract_pages(buf.getvalue(), "png"))


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_legacy_office_via_libreoffice():
    # LibreOffice reads an RTF body saved as .doc (it detects by content), which exercises the conversion path
    t = text_of(extract_pages(rb"{\rtf1\ansi Legacy Word file\par}", "doc"))
    assert "Legacy Word file" in t
